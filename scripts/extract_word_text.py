from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document


def clean_text(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\u3000", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def extract_docx(path: Path) -> dict[str, object]:
    try:
        document = Document(path)
        blocks: list[str] = []

        for paragraph in document.paragraphs:
            text = clean_text(paragraph.text)
            if text:
                blocks.append(text)

        for table_index, table in enumerate(document.tables, start=1):
            blocks.append(f"[[TABLE {table_index}]]")
            for row in table.rows:
                cells = [clean_text(cell.text) for cell in row.cells]
                if any(cells):
                    blocks.append("\t".join(cells))

        text = clean_text("\n".join(blocks))
        return {
            "source": path.name,
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "characters": len(text),
            "method": "python-docx",
            "text": text,
        }
    except Exception as error:
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        blocks = []
        for paragraph in root.iter(f"{namespace}p"):
            pieces = [node.text or "" for node in paragraph.iter(f"{namespace}t")]
            text = clean_text("".join(pieces))
            if text:
                blocks.append(text)
        text = clean_text("\n".join(blocks))
        return {
            "source": path.name,
            "paragraphs": len(blocks),
            "tables": None,
            "characters": len(text),
            "method": "ooxml-fallback",
            "warning": str(error),
            "text": text,
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    summary: list[dict[str, object]] = []
    for path in args.inputs:
        record = extract_docx(path)
        summary.append({key: value for key, value in record.items() if key != "text"})
        output = args.output_dir / f"{path.stem}.txt"
        output.write_text(str(record["text"]), encoding="utf-8")

    (args.output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
